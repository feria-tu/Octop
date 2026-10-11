"""Restrict managed assistants to private files and explicitly assigned knowledge."""

from __future__ import annotations

import io
from collections.abc import Awaitable, Callable
from pathlib import PurePosixPath
from typing import Any

from langchain.agents.middleware import AgentMiddleware
from langchain.agents.middleware.types import ModelRequest, ModelResponse
from langchain_core.messages import ToolMessage
from langchain_core.tools import StructuredTool
from langgraph.config import get_config
from langgraph.prebuilt.tool_node import ToolCallRequest
from langgraph.types import Command

from octop.infra.errors import ErrorCode, OctopError
from octop.infra.users.enterprise import SAFE_TOOLS, db_call


def extract_document_text(data: bytes, suffix: str) -> str:
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    if suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        return "\n".join(
            [
                *(p.text for p in document.paragraphs),
                *(
                    "\t".join(c.text for c in row.cells)
                    for table in document.tables
                    for row in table.rows
                ),
            ]
        )
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            return "\n".join(
                "\t".join(str(v) if v is not None else "" for v in row)
                for sheet in workbook
                for row in sheet.iter_rows(values_only=True)
            )
        finally:
            workbook.close()
    if suffix == ".pptx":
        from pptx import Presentation

        return "\n".join(
            str(shape.text)
            for slide in Presentation(io.BytesIO(data)).slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
    return data.decode("utf-8-sig")


def build_document_tool(workspace: Any) -> StructuredTool:
    async def read_document(path: str) -> str:
        """Extract text from a private uploaded PDF, DOCX, XLSX, PPTX or UTF-8 file."""
        normalized = PurePosixPath(path.replace("\\", "/").lstrip("/"))
        if (
            not normalized.parts
            or normalized.parts[0] not in {"inbound", "artifacts"}
            or ".." in normalized.parts
        ):
            raise OctopError.localized(ErrorCode.FORBIDDEN)
        from octop.infra.agents.workspace.enterprise import managed_content_path

        checked = await db_call(managed_content_path, workspace.workspace_dir, str(normalized))
        data = await workspace.adownload_bytes(checked)
        if data is None:
            raise OctopError.localized(ErrorCode.NOT_FOUND)
        return (await db_call(extract_document_text, data, normalized.suffix.lower()))[:100000]

    return StructuredTool.from_function(
        coroutine=read_document,
        name="read_document",
        description="Read text from an uploaded document in inbound/ or artifacts/. Supports PDF, DOCX, XLSX, PPTX and UTF-8 text; never executes commands.",
    )


class EnterpriseToolPolicyMiddleware(AgentMiddleware[Any, Any]):
    def __init__(self, repos: Any, agent_id: str, workspace: Any) -> None:
        self.repos = repos
        self.agent_id = agent_id
        self.workspace = workspace

    async def validate(self) -> None:
        configurable = get_config().get("configurable") or {}
        uid = configurable.get("user")
        user = await db_call(self.repos.user_repo.get, int(uid)) if uid is not None else None
        row = await db_call(self.repos.agent_repo.get, self.agent_id)
        if user is None or user.disabled or row is None:
            raise OctopError.localized(ErrorCode.FORBIDDEN)
        if row.enterprise_kind == "personal":
            if row.user_id != user.id and user.role != "admin":
                raise OctopError.localized(ErrorCode.FORBIDDEN)
            return
        records = await db_call(self.repos.enterprise_repo.runtimes, user_id=user.id)
        record = next((r for r in records if r.runtime_id == self.agent_id), None)
        definition = (
            await db_call(self.repos.agent_repo.get, record.definition_id) if record else None
        )
        if (
            record is None
            or definition is None
            or not definition.enabled
            or user.department_id != definition.department_id
            or record.definition_version != definition.definition_version
        ):
            raise OctopError.localized(ErrorCode.FORBIDDEN)

    async def awrap_model_call(
        self,
        request: ModelRequest[Any],
        handler: Callable[[ModelRequest[Any]], Awaitable[ModelResponse[Any]]],
    ) -> ModelResponse[Any]:
        await self.validate()
        allowed = SAFE_TOOLS | {"read_document"}
        tools = [
            tool
            for tool in request.tools
            if (tool.get("name") if isinstance(tool, dict) else getattr(tool, "name", ""))
            in allowed
        ]
        return await handler(request.override(tools=tools))

    async def awrap_tool_call(
        self,
        request: ToolCallRequest,
        handler: Callable[[ToolCallRequest], Awaitable[ToolMessage | Command[Any]]],
    ) -> ToolMessage | Command[Any]:
        await self.validate()
        if request.tool_call.get("name") not in SAFE_TOOLS | {"read_document"}:
            raise OctopError.localized(ErrorCode.FORBIDDEN)
        name = request.tool_call.get("name")
        if name in {
            "ls",
            "read_file",
            "glob",
            "grep",
            "write_file",
            "edit_file",
            "send_file_to_user",
            "read_document",
        }:
            from octop.infra.agents.workspace.enterprise import managed_content_path

            args = dict(request.tool_call.get("args") or {})
            path_keys = {"path", "file_path", "directory", "folder", "source", "output_path"}
            checked = False
            for key in path_keys & args.keys():
                if isinstance(args[key], str):
                    args[key] = await db_call(
                        managed_content_path, self.workspace.workspace_dir, args[key]
                    )
                    checked = True
            if name in {"ls", "glob", "grep"} and not checked:
                args["path"] = "inbound"
            request = request.override(tool_call={**request.tool_call, "args": args})
        return await handler(request)
